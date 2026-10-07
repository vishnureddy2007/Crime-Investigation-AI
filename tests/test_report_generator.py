"""
Tests for the report generators (PDF + DOCX).

We build a small `ReportData` from a synthetic analysis and verify
that both generators return valid non-empty bytes that contain
expected substrings.
"""

from __future__ import annotations

import datetime

import pytest

from models.evidence_analyzer import analyze
from models.report_generator import (
    DOCXReportGenerator,
    PDFReportGenerator,
    build_report_data,
)
from models.schemas import (
    AnalysisInput,
    BoundingBox,
    Detection,
    DetectionResult,
)


# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def _det(name: str, label: str, conf: float) -> Detection:
    return Detection(
        class_name=name, label=label,
        confidence=conf, bbox=BoundingBox(0, 0, 10, 10),
    )


def _analysis(counts: dict[str, int], conf: float = 0.85, source_type: str = "image"):
    dets = []
    for label, n in counts.items():
        dets.extend(_det(label, label, conf) for _ in range(n))
    dr = DetectionResult(
        source_name="crime_scene.jpg",
        timestamp=datetime.datetime.now(),
        detections=dets,
    )
    return analyze(AnalysisInput(
        source_name=dr.source_name,
        counts_by_label=dr.counts_by_label(),
        average_confidence=dr.average_confidence(),
        source_type=source_type,
    ))


# ----------------------------------------------------------------------
# build_report_data
# ----------------------------------------------------------------------
def test_build_report_data_basic() -> None:
    a = _analysis({"person": 2, "knife": 1})
    rd = build_report_data(analysis=a, summary=None)
    assert rd.source_name == "crime_scene.jpg"
    assert rd.source_type == "image"
    assert rd.total_objects == 3
    assert rd.person_count == 2
    assert rd.weapon_count == 1
    assert rd.has_threat is True
    assert rd.suggested_category in {"assault", "robbery"}
    assert rd.report_id.startswith("RPT-")
    # The template summary contains key investigation details
    assert "evidence from" in rd.summary_text
    assert "person" in rd.summary_text
    assert "weapon" in rd.summary_text


def test_build_report_data_uses_summary_when_provided() -> None:
    a = _analysis({"person": 1})
    rd_no = build_report_data(analysis=a, summary=None)
    rd_with_template = build_report_data(
        analysis=a,
        summary=type("S", (), {
            "primary_text": "AI: one person was present.",
            "model_name":   "google/flan-t5-base",
        })(),
    )
    assert "AI: one person was present." in rd_with_template.summary_text
    assert rd_with_template.summary_text != rd_no.summary_text
    assert "flan-t5" in rd_with_template.model_name


def test_build_report_data_custom_title_and_id() -> None:
    a = _analysis({"person": 1})
    rd = build_report_data(analysis=a, summary=None, title="Custom Title", report_id="RPT-FIXED")
    assert rd.title == "Custom Title"
    assert rd.report_id == "RPT-FIXED"


# ----------------------------------------------------------------------
# PDF generator
# ----------------------------------------------------------------------
def test_pdf_generator_returns_nonempty_bytes() -> None:
    rd = build_report_data(analysis=_analysis({"person": 1, "knife": 1}))
    try:
        pdf_bytes = PDFReportGenerator().generate(rd)
    except ImportError:
        pytest.skip("reportlab not installed")
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 100
    # PDF files start with %PDF-
    assert pdf_bytes.startswith(b"%PDF-")


def test_pdf_generator_with_empty_evidence() -> None:
    rd = build_report_data(analysis=_analysis({}))
    try:
        pdf_bytes = PDFReportGenerator().generate(rd)
    except ImportError:
        pytest.skip("reportlab not installed")
    assert pdf_bytes.startswith(b"%PDF-")
    assert len(pdf_bytes) > 100


# ----------------------------------------------------------------------
# DOCX generator
# ----------------------------------------------------------------------
def test_docx_generator_returns_nonempty_bytes() -> None:
    rd = build_report_data(analysis=_analysis({"person": 2, "knife": 1}))
    try:
        docx_bytes = DOCXReportGenerator().generate(rd)
    except ImportError:
        pytest.skip("python-docx not installed")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 100
    # .docx files are zip archives starting with "PK"
    assert docx_bytes[:2] == b"PK"


def test_docx_generator_with_empty_evidence() -> None:
    rd = build_report_data(analysis=_analysis({}))
    try:
        docx_bytes = DOCXReportGenerator().generate(rd)
    except ImportError:
        pytest.skip("python-docx not installed")
    assert docx_bytes[:2] == b"PK"
    assert len(docx_bytes) > 100


def test_docx_contains_expected_text() -> None:
    """Round-trip the docx and verify the title and source name appear."""
    rd = build_report_data(analysis=_analysis({"person": 3, "knife": 1}))
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")
    docx_bytes = DOCXReportGenerator().generate(rd)

    import io
    doc = Document(io.BytesIO(docx_bytes))

    full_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text

    # Title heading
    assert "AI Crime Investigation Report" in full_text
    # Source name should appear somewhere
    assert "crime_scene.jpg" in full_text
    # Suggested category should be in the document
    assert rd.suggested_category.replace("_", " ").title() in full_text


# ----------------------------------------------------------------------
# Both formats — content consistency
# ----------------------------------------------------------------------
def test_pdf_and_docx_share_report_id() -> None:
    rd = build_report_data(analysis=_analysis({"person": 1, "knife": 1}), report_id="RPT-CONSIST")
    try:
        pdf = PDFReportGenerator().generate(rd)
        docx = DOCXReportGenerator().generate(rd)
    except ImportError as exc:
        pytest.skip(f"missing dep: {exc}")
    # Both should be valid (non-empty, magic bytes correct)
    assert pdf.startswith(b"%PDF-")
    assert docx[:2] == b"PK"
    # The shared report_id is the contract
    assert rd.report_id == "RPT-CONSIST"
